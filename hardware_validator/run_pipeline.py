"""
Run script for Hardware Validation on Real Project Files.
Optimized to use a SINGLE LLM call with all rules in context.
"""
import sys
import os
import csv
import json
import hashlib

# Ensure project root is in path
project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

from hardware_validator.node import hardware_evaluator_node
from hardware_validator.parser import HardwareExportParser
from hardware_validator.schema import ValidationReport, ValidationIssue


def get_set_hash(hw_set):
    """Generates a unique hash for the hardware content of a set."""
    fingerprint = {
        "context": {
            "width": hw_set.opening_context.width_inches,
            "height": hw_set.opening_context.height_inches,
            "exterior": hw_set.opening_context.is_exterior,
            "fire": hw_set.opening_context.fire_rated
        },
        "components": [
            {"qty": c.quantity, "desc": c.description} 
            for c in hw_set.components
        ]
    }
    return hashlib.md5(json.dumps(fingerprint, sort_keys=True).encode()).hexdigest()


def run_validation_pipeline(excel_path: str):
    print(f"\n{'='*80}")
    print(f"HARDWARE VALIDATION PIPELINE (Optimized Single-Call)")
    print(f"Target: {os.path.basename(excel_path)}")
    print(f"{'='*80}")
    
    if not os.path.exists(excel_path):
        print(f"❌ Error: File not found at {excel_path}")
        return

    base_name = os.path.splitext(os.path.basename(excel_path))[0]
    
    # 1. Smart Sheet Detection
    print("\n[Phase 1] Scanning for Schedule Sheets...")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        
        # Look for Door Schedule sheets specifically
        schedule_sheets = [s for s in wb.sheetnames if "Door Schedule" in s]
        if not schedule_sheets:
            # Fallback to any schedule-like sheets
            schedule_sheets = [s for s in wb.sheetnames if "Schedule" in s or "Matl" in s]
        
        if not schedule_sheets:
            print("  ⚠ No schedule sheets found.")
            return
            
        print(f"  Found {len(schedule_sheets)} schedule sheet(s): {schedule_sheets}")
        
        total_issues_found = 0
        
        for sheet_name in schedule_sheets:
            print(f"\n{'─'*60}")
            print(f"[Phase 2] Processing: {sheet_name}")
            print(f"{'─'*60}")
            
            target_sheet = wb[sheet_name]
            output_dir = os.path.dirname(__file__)
            csv_extract_path = os.path.join(output_dir, f"{base_name}_{sheet_name.replace(' ', '_').replace('(', '').replace(')', '')}_extract.csv")
            annotated_path = os.path.join(output_dir, f"{base_name}_{sheet_name.replace(' ', '_').replace('(', '').replace(')', '')}_annotated.csv")
            
            # Extract to CSV
            with open(csv_extract_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                for row in target_sheet.iter_rows(values_only=True):
                    clean_row = [str(c) if c is not None else "" for c in row]
                    writer.writerow(clean_row)
            
            # Parse openings
            all_hw_sets = HardwareExportParser.parse_csv(csv_extract_path)
            if not all_hw_sets:
                print(f"  (Skipping: No valid openings found)")
                continue
                
            print(f"  Parsed {len(all_hw_sets)} openings")

            # Deduplicate by hardware configuration
            unique_sets_map = {}
            set_to_openings_map = {}
            
            for hw_set in all_hw_sets:
                h = get_set_hash(hw_set)
                if h not in unique_sets_map:
                    unique_sets_map[h] = hw_set
                    set_to_openings_map[h] = []
                set_to_openings_map[h].append(hw_set.set_id)
                
            unique_sets = list(unique_sets_map.values())
            print(f"  Consolidated to {len(unique_sets)} unique configurations")

            # Single validation call for all unique sets
            print(f"\n[Phase 3] AI Validation (Single Call)...")
            state = {"hardware_sets": unique_sets}
            result = hardware_evaluator_node(state)
            
            # Fan-out results to all openings
            final_reports = []
            result_map = {r["set_id"]: r for r in result["validation_reports"]}
            
            for h, hw_set in unique_sets_map.items():
                rep_result = result_map.get(hw_set.set_id)
                if not rep_result:
                    continue
                
                for opening_id in set_to_openings_map[h]:
                    clean_issues = []
                    for i_data in rep_result["issues"]:
                        clean_issues.append(ValidationIssue(
                            level=i_data['level'],
                            category=i_data['category'],
                            message=i_data['message'],
                            component_index=i_data.get('component_index'),
                            fix_suggestion=i_data.get('fix_suggestion'),
                            source_line=i_data.get('source_line')
                        ))
                    
                    final_reports.append(ValidationReport(
                        set_id=opening_id,
                        issues=clean_issues,
                        score=rep_result["score"]
                    ))

            # Write annotated CSV
            HardwareExportParser.write_annotated_csv(csv_extract_path, annotated_path, final_reports)
            
            # Summary - only count ERROR and WARNING level issues (not INFO)
            openings_with_real_issues = [
                r for r in final_reports 
                if any(i.level in ("ERROR", "WARNING") for i in r.issues)
            ]
            sheet_issues = len(openings_with_real_issues)
            total_issues_found += sheet_issues
            
            print(f"\n[Phase 4] Results Summary")
            if sheet_issues > 0:
                print(f"  ⚠ Found issues in {sheet_issues} opening(s)")
                print(f"\n  {'─'*50}")
                
                # Show unique issues (not repeated for each opening)
                shown_issues = set()
                for r_data in result["validation_reports"]:
                    # Only show if there are ERROR or WARNING level issues
                    real_issues = [i for i in r_data["issues"] if i['level'] in ("ERROR", "WARNING")]
                    if real_issues:
                        # Find how many openings share this configuration
                        affected_count = 0
                        for h, s in unique_sets_map.items():
                            if s.set_id == r_data['set_id']:
                                affected_count = len(set_to_openings_map[h])
                                break
                        
                        print(f"\n  📋 Configuration: {r_data['set_id']} (affects {affected_count} opening(s))")
                        print(f"     Score: {r_data['score']}/100")
                        
                        for issue in real_issues:
                            issue_key = f"{issue['category']}:{issue['message']}"
                            if issue_key not in shown_issues:
                                shown_issues.add(issue_key)
                                level_icon = "❌" if issue['level'] == "ERROR" else "⚠️"
                                print(f"     {level_icon} [{issue['level']}] {issue['message']}")
                                if issue.get('fix_suggestion'):
                                    print(f"        💡 Fix: {issue['fix_suggestion']}")
            else:
                print("  ✅ No issues found - all configurations pass validation")

            print(f"\n  📄 Report saved: {os.path.basename(annotated_path)}")
        
        print(f"\n{'='*80}")
        if total_issues_found > 0:
            print(f"VALIDATION COMPLETE: {total_issues_found} opening(s) with issues across all sheets")
        else:
            print(f"VALIDATION COMPLETE: All openings pass validation ✅")
        print(f"{'='*80}\n")
        
    except Exception as e:
        print(f"  ❌ Pipeline error: {e}")
        import traceback
        traceback.print_exc()
        return


if __name__ == "__main__":
    default_path = os.path.abspath(os.path.join(os.getcwd(), "../../comsense_projects/000/TxDMV Camp Hubbard - The Door Company.xlsm"))
    if len(sys.argv) > 1:
        default_path = sys.argv[1]
        
    run_validation_pipeline(default_path)
