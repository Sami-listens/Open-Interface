import sys
import os
# Force import from site-packages if conda environment activation is flaky in this shell context
sys.path.append('/opt/anaconda3/envs/door_company/lib/python3.11/site-packages')

import openpyxl

project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))

def analyze_excel_structure():
    # Adjust file path to point to the known location
    file_path = os.path.join(project_root, "Project Example - THM WWTP", "1. Estimates", "00-Pricing", "THM WWTP - Project Tool - 8C.xlsm")
    print(f"Analyzing: {file_path}")
    
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        print("Sheet Names:", wb.sheetnames)
        
        # Since no explicit "Hardware Schedule" was found, let's dump headers of relevant-looking sheets
        # 'Estimate - Material' might contain hardware
        # 'Door Schedule' definitely contains openings
        
        for sheet_name in ['Estimate - Material', 'Door Schedule', 'Door Schedule(2)']:
            if sheet_name in wb.sheetnames:
                print(f"\n--- Content of {sheet_name} (First 10 rows) ---")
                ws = wb[sheet_name]
                for i, row in enumerate(ws.iter_rows(max_row=10, values_only=True)):
                    clean_row = [str(cell)[:30] if cell is not None else "" for cell in row]
                    print(f"Row {i+1}: {clean_row}")

    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    analyze_excel_structure()

